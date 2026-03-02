import io
import logging
from datetime import date, datetime, timedelta

from django.contrib.auth import login, logout, authenticate
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.db import transaction
from django.db.models import Count, Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.utils import timezone
from django.views.decorators.http import require_POST

from .decorators import admin_required, staff_or_above, coordinator_or_admin
from .forms import (
    LoginForm, UserCreateForm, UserEditForm, DepartmentForm,
    StudentForm, AttendanceFilterForm, EditRequestForm, ReportFilterForm
)
from .models import (
    CustomUser, Department, Student, Period,
    AttendanceSession, AttendanceRecord, EditRequest
)
from .sms import send_absent_sms

logger = logging.getLogger(__name__)


# ─── Auth ────────────────────────────────────────────────────────────────────

def login_view(request):
    if request.user.is_authenticated:
        return redirect('dashboard')
    form = LoginForm(request, data=request.POST or None)
    if request.method == 'POST' and form.is_valid():
        login(request, form.get_user())
        messages.success(request, f"Welcome, {form.get_user().get_full_name() or form.get_user().username}!")
        return redirect(request.GET.get('next', 'dashboard'))
    return render(request, 'attendance/login.html', {'form': form})


@login_required
def logout_view(request):
    logout(request)
    messages.info(request, "You have been logged out.")
    return redirect('login')


# ─── Dashboard ───────────────────────────────────────────────────────────────

@login_required
def dashboard(request):
    today = date.today()
    context = {'today': today}

    if request.user.role == 'admin':
        total_students = Student.objects.filter(is_active=True).count()
        total_departments = Department.objects.count()
        today_sessions = AttendanceSession.objects.filter(date=today, is_submitted=True)
        total_present_today = AttendanceRecord.objects.filter(
            session__date=today, session__is_submitted=True, status='present'
        ).count()
        total_absent_today = AttendanceRecord.objects.filter(
            session__date=today, session__is_submitted=True, status='absent'
        ).count()
        pending_edit_requests = EditRequest.objects.filter(status='pending').count()
        dept_summary = []
        for dept in Department.objects.all():
            sessions = AttendanceSession.objects.filter(date=today, department=dept, is_submitted=True)
            present = AttendanceRecord.objects.filter(
                session__in=sessions, status='present'
            ).count()
            absent = AttendanceRecord.objects.filter(
                session__in=sessions, status='absent'
            ).count()
            dept_summary.append({
                'dept': dept, 'present': present, 'absent': absent,
                'sessions': sessions.count()
            })
        context.update({
            'total_students': total_students,
            'total_departments': total_departments,
            'today_sessions': today_sessions.count(),
            'total_present_today': total_present_today,
            'total_absent_today': total_absent_today,
            'pending_edit_requests': pending_edit_requests,
            'dept_summary': dept_summary,
        })
        return render(request, 'attendance/dashboard_admin.html', context)

    elif request.user.role == 'staff':
        dept = request.user.department
        my_sessions_today = AttendanceSession.objects.filter(
            date=today, department=dept, marked_by=request.user
        )
        submitted_periods = my_sessions_today.filter(is_submitted=True).values_list(
            'period__period_number', flat=True
        )
        periods = Period.objects.all()
        context.update({
            'dept': dept,
            'my_sessions_today': my_sessions_today,
            'submitted_periods': list(submitted_periods),
            'periods': periods,
        })
        return render(request, 'attendance/dashboard_staff.html', context)

    elif request.user.role == 'coordinator':
        pending_requests = EditRequest.objects.filter(status='pending').select_related(
            'record__student', 'record__session', 'requested_by'
        )
        today_sessions = AttendanceSession.objects.filter(date=today, is_submitted=True).count()
        total_absent = AttendanceRecord.objects.filter(
            session__date=today, session__is_submitted=True, status='absent'
        ).count()
        context.update({
            'pending_requests': pending_requests[:10],
            'pending_count': pending_requests.count(),
            'today_sessions': today_sessions,
            'total_absent': total_absent,
        })
        return render(request, 'attendance/dashboard_coordinator.html', context)

    return redirect('login')


# ─── Department Management ────────────────────────────────────────────────────

@login_required
@admin_required
def department_list(request):
    departments = Department.objects.annotate(student_count=Count('students'))
    return render(request, 'attendance/department_list.html', {'departments': departments})


@login_required
@admin_required
def department_create(request):
    form = DepartmentForm(request.POST or None)
    if request.method == 'POST' and form.is_valid():
        form.save()
        messages.success(request, "Department created successfully.")
        return redirect('department_list')
    return render(request, 'attendance/department_form.html', {'form': form, 'title': 'Add Department'})


@login_required
@admin_required
def department_edit(request, pk):
    dept = get_object_or_404(Department, pk=pk)
    form = DepartmentForm(request.POST or None, instance=dept)
    if request.method == 'POST' and form.is_valid():
        form.save()
        messages.success(request, "Department updated successfully.")
        return redirect('department_list')
    return render(request, 'attendance/department_form.html', {'form': form, 'title': 'Edit Department'})


# ─── User Management ─────────────────────────────────────────────────────────

@login_required
@admin_required
def user_list(request):
    users = CustomUser.objects.exclude(pk=request.user.pk).select_related('department')
    return render(request, 'attendance/user_list.html', {'users': users})


@login_required
@admin_required
def user_create(request):
    form = UserCreateForm(request.POST or None)
    if request.method == 'POST' and form.is_valid():
        form.save()
        messages.success(request, "User created successfully.")
        return redirect('user_list')
    return render(request, 'attendance/user_form.html', {'form': form, 'title': 'Add User'})


@login_required
@admin_required
def user_edit(request, pk):
    user = get_object_or_404(CustomUser, pk=pk)
    form = UserEditForm(request.POST or None, instance=user)
    if request.method == 'POST' and form.is_valid():
        form.save()
        messages.success(request, "User updated successfully.")
        return redirect('user_list')
    return render(request, 'attendance/user_form.html', {'form': form, 'title': 'Edit User'})


# ─── Student Management ───────────────────────────────────────────────────────

@login_required
@staff_or_above
def student_list(request):
    dept_filter = request.GET.get('department')
    students = Student.objects.select_related('department')
    if request.user.role == 'staff' and request.user.department:
        students = students.filter(department=request.user.department)
    elif dept_filter:
        students = students.filter(department_id=dept_filter)
    departments = Department.objects.all()
    return render(request, 'attendance/student_list.html', {
        'students': students,
        'departments': departments,
        'selected_dept': dept_filter,
    })


@login_required
@admin_required
def student_create(request):
    form = StudentForm(request.POST or None)
    if request.method == 'POST' and form.is_valid():
        form.save()
        messages.success(request, "Student added successfully.")
        return redirect('student_list')
    return render(request, 'attendance/student_form.html', {'form': form, 'title': 'Add Student'})


@login_required
@admin_required
def student_edit(request, pk):
    student = get_object_or_404(Student, pk=pk)
    form = StudentForm(request.POST or None, instance=student)
    if request.method == 'POST' and form.is_valid():
        form.save()
        messages.success(request, "Student updated successfully.")
        return redirect('student_list')
    return render(request, 'attendance/student_form.html', {'form': form, 'title': 'Edit Student'})


@login_required
@staff_or_above
def student_detail(request, pk):
    student = get_object_or_404(Student, pk=pk)
    records = AttendanceRecord.objects.filter(student=student).select_related(
        'session__period', 'session__department'
    ).order_by('-session__date', 'session__period__period_number')
    percentage = student.attendance_percentage()
    return render(request, 'attendance/student_detail.html', {
        'student': student,
        'records': records[:50],
        'percentage': percentage,
    })


# ─── Attendance Marking ───────────────────────────────────────────────────────

@login_required
@staff_or_above
def mark_attendance(request):
    today = date.today()
    periods = Period.objects.all()
    departments = Department.objects.all()

    # Staff see only their department
    if request.user.role == 'staff':
        departments = departments.filter(pk=request.user.department.pk) if request.user.department else departments.none()

    selected_date = request.GET.get('date', str(today))
    selected_period_id = request.GET.get('period')
    selected_dept_id = request.GET.get('department')

    if request.user.role == 'staff' and request.user.department:
        selected_dept_id = selected_dept_id or str(request.user.department.pk)

    context = {
        'periods': periods,
        'departments': departments,
        'selected_date': selected_date,
        'selected_period_id': selected_period_id,
        'selected_dept_id': selected_dept_id,
        'today': str(today),
    }

    if selected_period_id and selected_dept_id and selected_date:
        period = get_object_or_404(Period, pk=selected_period_id)
        dept = get_object_or_404(Department, pk=selected_dept_id)

        # Check for existing submitted session
        existing_session = AttendanceSession.objects.filter(
            date=selected_date, period=period, department=dept
        ).first()

        if existing_session and existing_session.is_submitted and request.user.role == 'staff':
            messages.warning(request, "Attendance for this session has already been submitted and cannot be edited.")
            context['session'] = existing_session
            context['records'] = existing_session.records.select_related('student').order_by('student__register_number')
            context['is_submitted'] = True
            return render(request, 'attendance/mark_attendance.html', context)

        students = Student.objects.filter(department=dept, is_active=True).order_by('register_number')
        context.update({
            'period': period, 'dept': dept, 'students': students,
            'session': existing_session,
            'is_submitted': existing_session.is_submitted if existing_session else False,
        })

        if existing_session:
            existing_records = {r.student_id: r for r in existing_session.records.all()}
            context['existing_records'] = existing_records

        if request.method == 'POST':
            action = request.POST.get('action', 'save')
            with transaction.atomic():
                session, created = AttendanceSession.objects.get_or_create(
                    date=selected_date, period=period, department=dept,
                    defaults={'marked_by': request.user}
                )
                if session.is_submitted and request.user.role == 'staff':
                    messages.error(request, "Cannot edit a submitted session.")
                    return redirect('mark_attendance')

                for student in students:
                    status = request.POST.get(f'status_{student.pk}', 'present')
                    record, _ = AttendanceRecord.objects.get_or_create(
                        session=session, student=student,
                        defaults={'status': status}
                    )
                    if not _:
                        record.status = status
                        record.save()

                    # Send SMS if absent and not yet sent
                    if status == 'absent' and not record.sms_sent:
                        send_absent_sms(student, session)
                        record.sms_sent = True
                        record.save(update_fields=['sms_sent'])

                if action == 'submit':
                    session.is_submitted = True
                    session.submitted_at = timezone.now()
                    session.save()
                    messages.success(request, f"Attendance submitted for {dept.name} | Period {period.period_number}.")
                    return redirect('mark_attendance')
                else:
                    messages.success(request, "Attendance saved (draft).")
                    return redirect(
                        f'/attendance/mark/?date={selected_date}&period={selected_period_id}&department={selected_dept_id}'
                    )

    return render(request, 'attendance/mark_attendance.html', context)


# ─── Edit Requests ────────────────────────────────────────────────────────────

@login_required
@coordinator_or_admin
def edit_request_list(request):
    if request.user.role == 'coordinator':
        requests = EditRequest.objects.filter(requested_by=request.user).select_related(
            'record__student', 'record__session__period', 'record__session__department'
        )
        pending = EditRequest.objects.filter(status='pending').select_related(
            'record__student', 'record__session__period', 'record__session__department', 'requested_by'
        )
    else:
        requests = EditRequest.objects.all().select_related(
            'record__student', 'record__session__period', 'record__session__department', 'requested_by'
        )
        pending = requests.filter(status='pending')
    return render(request, 'attendance/edit_request_list.html', {
        'requests': requests,
        'pending': pending,
    })


@login_required
@coordinator_or_admin
def create_edit_request(request, record_id):
    record = get_object_or_404(AttendanceRecord, pk=record_id)
    if not record.session.is_submitted:
        messages.error(request, "Can only request edits for submitted sessions.")
        return redirect('mark_attendance')

    form = EditRequestForm(request.POST or None)
    if request.method == 'POST' and form.is_valid():
        er = form.save(commit=False)
        er.session = record.session
        er.record = record
        er.requested_by = request.user
        er.old_status = record.status
        er.save()
        messages.success(request, "Edit request submitted successfully.")
        return redirect('edit_request_list')
    return render(request, 'attendance/edit_request_form.html', {
        'form': form, 'record': record
    })


@login_required
@admin_required
def review_edit_request(request, pk):
    er = get_object_or_404(EditRequest, pk=pk, status='pending')
    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'approve':
            with transaction.atomic():
                er.record.status = er.new_status
                er.record.save()
                # Send SMS if newly marked absent
                if er.new_status == 'absent' and not er.record.sms_sent:
                    send_absent_sms(er.record.student, er.record.session)
                    er.record.sms_sent = True
                    er.record.save(update_fields=['sms_sent'])
                er.status = 'approved'
                er.approved_by = request.user
                er.reviewed_at = timezone.now()
                er.save()
            messages.success(request, "Edit request approved and attendance updated.")
        elif action == 'reject':
            er.status = 'rejected'
            er.approved_by = request.user
            er.reviewed_at = timezone.now()
            er.save()
            messages.info(request, "Edit request rejected.")
        return redirect('edit_request_list')
    return render(request, 'attendance/review_edit_request.html', {'er': er})


# ─── Reports ──────────────────────────────────────────────────────────────────

@login_required
@staff_or_above
def attendance_report(request):
    today = date.today()
    form = ReportFilterForm(request.GET or None)
    report_data = []
    selected_dept = None
    start_date = None
    end_date = None

    if form.is_valid() or request.GET:
        dept_id = request.GET.get('department')
        month = request.GET.get('month')
        year = request.GET.get('year')
        start_date = request.GET.get('start_date')
        end_date = request.GET.get('end_date')

        if request.user.role == 'staff' and request.user.department:
            dept_id = str(request.user.department.pk)

        students = Student.objects.filter(is_active=True).select_related('department')
        if dept_id:
            students = students.filter(department_id=dept_id)
            try:
                selected_dept = Department.objects.get(pk=dept_id)
            except Department.DoesNotExist:
                pass

        if start_date and end_date:
            try:
                start_date = datetime.strptime(start_date, '%Y-%m-%d').date()
                end_date = datetime.strptime(end_date, '%Y-%m-%d').date()
            except ValueError:
                start_date = end_date = None
        elif month and year:
            import calendar
            year = int(year)
            month = int(month)
            start_date = date(year, month, 1)
            end_date = date(year, month, calendar.monthrange(year, month)[1])
        else:
            start_date = today.replace(day=1)
            end_date = today

        for student in students:
            pct = student.attendance_percentage(start_date=start_date, end_date=end_date)
            records = AttendanceRecord.objects.filter(
                student=student,
                session__date__gte=start_date,
                session__date__lte=end_date
            )
            total = records.count()
            present = records.filter(status='present').count()
            absent = records.filter(status='absent').count()
            il = records.filter(status='informed_leave').count()
            report_data.append({
                'student': student,
                'total': total,
                'present': present,
                'absent': absent,
                'informed_leave': il,
                'percentage': pct,
            })

    return render(request, 'attendance/attendance_report.html', {
        'form': form,
        'report_data': report_data,
        'selected_dept': selected_dept,
        'start_date': start_date,
        'end_date': end_date,
    })


@login_required
@staff_or_above
def daily_summary(request):
    selected_date = request.GET.get('date', str(date.today()))
    try:
        selected_date = datetime.strptime(selected_date, '%Y-%m-%d').date()
    except ValueError:
        selected_date = date.today()

    sessions = AttendanceSession.objects.filter(
        date=selected_date, is_submitted=True
    ).select_related('period', 'department', 'marked_by')

    if request.user.role == 'staff' and request.user.department:
        sessions = sessions.filter(department=request.user.department)

    summary = []
    for session in sessions:
        records = session.records.all()
        total = records.count()
        present = records.filter(status='present').count()
        absent = records.filter(status='absent').count()
        il = records.filter(status='informed_leave').count()
        summary.append({
            'session': session,
            'total': total,
            'present': present,
            'absent': absent,
            'informed_leave': il,
            'percentage': round((present / total * 100), 2) if total else 0,
        })

    return render(request, 'attendance/daily_summary.html', {
        'summary': summary,
        'selected_date': selected_date,
    })


# ─── Export ───────────────────────────────────────────────────────────────────

@login_required
@staff_or_above
def export_excel(request):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    dept_id = request.GET.get('department')
    start_date_str = request.GET.get('start_date')
    end_date_str = request.GET.get('end_date')

    students = Student.objects.filter(is_active=True).select_related('department')
    if dept_id:
        students = students.filter(department_id=dept_id)
    if request.user.role == 'staff' and request.user.department:
        students = students.filter(department=request.user.department)

    start_date = end_date = None
    if start_date_str:
        try:
            start_date = datetime.strptime(start_date_str, '%Y-%m-%d').date()
        except ValueError:
            pass
    if end_date_str:
        try:
            end_date = datetime.strptime(end_date_str, '%Y-%m-%d').date()
        except ValueError:
            pass

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance Report"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    headers = ['Reg. No.', 'Student Name', 'Department', 'Total Sessions', 'Present', 'Absent', 'Informed Leave', 'Attendance %']

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    for row, student in enumerate(students, 2):
        records = AttendanceRecord.objects.filter(student=student)
        if start_date:
            records = records.filter(session__date__gte=start_date)
        if end_date:
            records = records.filter(session__date__lte=end_date)
        total = records.count()
        present = records.filter(status='present').count()
        absent = records.filter(status='absent').count()
        il = records.filter(status='informed_leave').count()
        pct = round((present / total * 100), 2) if total else 0

        ws.append([
            student.register_number, student.name, student.department.name,
            total, present, absent, il, f"{pct}%"
        ])

        if pct < 75:
            red_fill = PatternFill(start_color="FFE0E0", end_color="FFE0E0", fill_type="solid")
            for col in range(1, 9):
                ws.cell(row=row, column=col).fill = red_fill

    for col in ws.columns:
        max_length = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_length + 4

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="attendance_report.xlsx"'
    wb.save(response)
    return response


@login_required
@staff_or_above
def export_pdf(request):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.units import inch

    dept_id = request.GET.get('department')
    start_date_str = request.GET.get('start_date')
    end_date_str = request.GET.get('end_date')

    students = Student.objects.filter(is_active=True).select_related('department')
    if dept_id:
        students = students.filter(department_id=dept_id)
    if request.user.role == 'staff' and request.user.department:
        students = students.filter(department=request.user.department)

    start_date = end_date = None
    if start_date_str:
        try:
            start_date = datetime.strptime(start_date_str, '%Y-%m-%d').date()
        except ValueError:
            pass
    if end_date_str:
        try:
            end_date = datetime.strptime(end_date_str, '%Y-%m-%d').date()
        except ValueError:
            pass

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4))
    styles = getSampleStyleSheet()
    elements = []

    title_style = ParagraphStyle('title', parent=styles['Heading1'], alignment=1, spaceAfter=12)
    elements.append(Paragraph("College Attendance Report", title_style))
    if start_date and end_date:
        elements.append(Paragraph(f"Period: {start_date} to {end_date}", styles['Normal']))
    elements.append(Spacer(1, 0.2 * inch))

    data = [['Reg. No.', 'Student Name', 'Department', 'Total', 'Present', 'Absent', 'Inf. Leave', 'Attendance %']]
    for student in students:
        records = AttendanceRecord.objects.filter(student=student)
        if start_date:
            records = records.filter(session__date__gte=start_date)
        if end_date:
            records = records.filter(session__date__lte=end_date)
        total = records.count()
        present = records.filter(status='present').count()
        absent = records.filter(status='absent').count()
        il = records.filter(status='informed_leave').count()
        pct = round((present / total * 100), 2) if total else 0
        data.append([
            student.register_number, student.name, student.department.name,
            total, present, absent, il, f"{pct}%"
        ])

    col_widths = [1.2*inch, 2*inch, 1.5*inch, 0.7*inch, 0.7*inch, 0.7*inch, 0.8*inch, 1*inch]
    table = Table(data, colWidths=col_widths)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4E79')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#EBF3FB')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('ROWHEIGHT', (0, 0), (-1, -1), 20),
    ]))

    elements.append(table)
    doc.build(elements)

    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="attendance_report.pdf"'
    return response


# ─── Attendance View Detail ───────────────────────────────────────────────────

@login_required
@staff_or_above
def session_detail(request, pk):
    session = get_object_or_404(AttendanceSession, pk=pk)
    if request.user.role == 'staff' and session.department != request.user.department:
        messages.error(request, "Access denied.")
        return redirect('dashboard')
    records = session.records.select_related('student').order_by('student__register_number')
    return render(request, 'attendance/session_detail.html', {
        'session': session, 'records': records
    })


@login_required
@staff_or_above
def sessions_list(request):
    sessions = AttendanceSession.objects.select_related('period', 'department', 'marked_by')
    if request.user.role == 'staff' and request.user.department:
        sessions = sessions.filter(department=request.user.department)
    sessions = sessions.order_by('-date', 'period__period_number')[:100]
    return render(request, 'attendance/sessions_list.html', {'sessions': sessions})
